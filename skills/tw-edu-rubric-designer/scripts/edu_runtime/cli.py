from __future__ import annotations

import argparse
import hashlib
import html
import json
import math
import re
import sys
import uuid
from pathlib import Path
from typing import Any

SKILLS = {
    "tw-edu-learning-evidence-analyzer":"docx", "tw-edu-material-reviewer":"docx",
    "tw-edu-lesson-plan-108": "docx", "tw-edu-curriculum-mapper": "xlsx",
    "tw-edu-exam-generator": "exam", "tw-edu-rubric-designer": "docx",
    "tw-edu-feedback-writer": "docx", "tw-edu-learning-portfolio": "docx",
    "tw-edu-classroom-culture": "docx", "tw-edu-differentiated": "docx",
    "tw-edu-formative-assessment": "docx", "tw-edu-interdisciplinary": "docx",
    "tw-edu-meeting-facilitator": "docx", "tw-edu-parent-communication": "docx",
    "tw-edu-pbl-designer": "docx", "tw-edu-school-document": "docx",
    "tw-edu-worksheet-creator": "docx", "tw-edu-anti-ai-assessment": "docx",
    "tw-edu-mini-app": "html", "tw-edu-research-viz": "png",
    "tw-edu-slides-creator": "pptx",
}

class InputError(Exception): pass

def _json(path: Path) -> Any:
    try: return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as e: raise InputError(f"cannot read JSON {path}: {e}") from e

def _hash(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(65536), b""): h.update(block)
    return h.hexdigest()

def _skill_dir() -> Path:
    # <skill>/scripts/edu_runtime/cli.py
    return Path(__file__).resolve().parents[2]

def _validate(data: dict, schema_path: Path) -> None:
    def check_values(value, location='$'):
        if isinstance(value, str) and not value.strip():
            raise InputError(f'empty text at {location}')
        if isinstance(value, float) and not math.isfinite(value):
            raise InputError(f'non-finite number at {location}')
        if isinstance(value, dict):
            for key, item in value.items(): check_values(item, f'{location}.{key}')
        elif isinstance(value, list):
            for index, item in enumerate(value): check_values(item, f'{location}[{index}]')
    check_values(data)
    try:
        import jsonschema
    except ImportError as e: raise InputError("jsonschema is required") from e
    schema = _json(schema_path)
    try: jsonschema.validate(data, schema)
    except jsonschema.ValidationError as e:
        location = ".".join(str(p) for p in e.absolute_path) or "$"
        raise InputError(f"validation failed at {location}: {e.message}") from e

def _semantic(skill: str, d: dict) -> None:
    c = d["content"]
    if skill == "tw-edu-learning-evidence-analyzer":
        from .learning_evidence import analyze
        try: analyze(c)
        except (ValueError, KeyError, TypeError) as exc: raise InputError(str(exc)) from exc
    if skill == "tw-edu-material-reviewer":
        materials={x['id']:x['text'] for x in c['materials']}
        if len(materials)!=len(c['materials']): raise InputError('duplicate material ID')
        if len({f['id'] for f in c['findings']})!=len(c['findings']): raise InputError('duplicate finding ID')
        for f in c['findings']:
            if f['material_id'] not in materials or f['quote'] not in materials[f['material_id']]: raise InputError('finding quote is not located in material')
            if f['status']=='unknown' and f['severity']!='unknown': raise InputError('unknown finding must retain unknown severity')
    def unique(items, field, label):
        values=[x[field] for x in items]
        if len(values)!=len(set(values)): raise InputError(f"duplicate {label}")
    if skill == "tw-edu-lesson-plan-108":
        unique(c["objectives"],"id","objective id"); unique(c["activities"],"id","activity id"); unique(c["assessments"],"id","assessment id")
        ids = {x["id"] for x in c["objectives"]}
        if sum(x["minutes"] for x in c["activities"]) != c["total_minutes"]: raise InputError("activity minutes must equal total_minutes")
        for group in (c["activities"], c["assessments"]):
            if any(not set(x["objective_ids"]).issubset(ids) for x in group): raise InputError("unknown objective_id")
        taught=set().union(*(set(x["objective_ids"]) for x in c["activities"])); assessed=set().union(*(set(x["objective_ids"]) for x in c["assessments"]))
        if taught != ids: raise InputError(f"every objective must be taught; uncovered: {sorted(ids-taught)}")
        if assessed != ids: raise InputError(f"every objective must be assessed; uncovered: {sorted(ids-assessed)}")
        for code in c.get("curriculum_codes", []):
            if not code.get("verified", False) and not code.get("verification_note"): raise InputError("unverified curriculum code requires verification_note")
    elif skill == "tw-edu-exam-generator":
        unique(c["questions"],"id","question id")
        if len(c["questions"]) != c["expected_question_count"]: raise InputError("actual question count mismatch")
        if sum(q["points"] for q in c["questions"]) != c["total_points"]: raise InputError("question points do not equal total_points")
        for q in c["questions"]:
            if q.get("options"): unique(q["options"],"id",f"option id in question {q['id']}")
            if q["type"] == "multiple_choice" and (q.get("answer") not in [o["id"] for o in q.get("options", [])]): raise InputError(f"question {q['id']} answer is not an option")
    elif skill == "tw-edu-rubric-designer":
        levels = c["levels"]
        unique(levels, 'id', 'level id')
        if c["type"] == "analytic":
            unique(c['dimensions'], 'id', 'dimension id')
            if any(set(x["descriptions"]) != {l["id"] for l in levels} for x in c["dimensions"]): raise InputError("each analytic dimension needs every level description")
            if sum(x["weight"] for x in c["dimensions"]) != c["total_points"]: raise InputError("dimension weights do not equal total_points")
        elif set(c["descriptions"]) != {l["id"] for l in levels}: raise InputError("holistic rubric needs every level description")
        elif max(l["score"] for l in levels) != c["total_points"]: raise InputError("holistic total_points must equal the highest level score")
    elif skill == "tw-edu-curriculum-mapper":
        for unit in c["units"]:
            for code in unit["codes"]:
                if not code["verified"] and not code.get("verification_note"): raise InputError("unverified curriculum code requires verification_note")
    elif skill == "tw-edu-research-viz" and c["type"] == "prisma":
        p = c["prisma"]
        if p["identified"] != p["duplicates_removed"] + p["screened"]: raise InputError("PRISMA identification counts are not conserved")
        if p["screened"] != p["screening_excluded"] + p["full_text_assessed"]: raise InputError("PRISMA screening counts are not conserved")
        if p["full_text_assessed"] != p["full_text_excluded"] + p["included"]: raise InputError("PRISMA eligibility counts are not conserved")
    elif skill == "tw-edu-mini-app" and c["mode"] == "quiz":
        unique(c["questions"],"id","question id")
        for q in c["questions"]:
            unique(q['options'], 'id', 'option id')
            if q["answer"] not in [o["id"] for o in q["options"]]: raise InputError(f"question {q['id']} answer is not an option")
    elif skill == "tw-edu-anti-ai-assessment":
        for item in c["items"]:
            if sum(x["score"] for x in item["dimensions"]) != item["total_score"]: raise InputError(f"item {item['id']} dimension scores do not equal total_score")
    elif skill == "tw-edu-slides-creator":
        ids=[x["id"] for x in c["slides"]]
        if ids != list(range(1,len(ids)+1)): raise InputError("slide ids must be unique and contiguous from 1")
        for slide in c["slides"]:
            if slide["type"]=="chart":
                n=len(slide["chart"]["categories"])
                if any(len(s["values"])!=n for s in slide["chart"]["series"]): raise InputError(f"slide {slide['id']} chart category/value lengths differ")

DOC_LABELS={
"tw-edu-feedback-writer":{"students":"學生觀察與回饋"},"tw-edu-learning-portfolio":{"records":"學習歷程紀錄"},
"tw-edu-classroom-culture":{"agreements":"班級共識","routines":"日常程序","response_plan":"事件回應計畫"},
"tw-edu-differentiated":{"shared_goal":"共同學習目標","learner_groups":"學習需求與支持","activities":"差異化活動"},
"tw-edu-formative-assessment":{"learning_target":"學習目標","checks":"評量檢核","response_rules":"依證據調整教學"},
"tw-edu-interdisciplinary":{"disciplines":"跨域學科","driving_question":"驅動問題","discipline_contributions":"各科貢獻","activities":"學習活動","product":"成果作品"},
"tw-edu-meeting-facilitator":{"participants":"與會人員","agenda":"議程","decisions":"決議","actions":"待辦追蹤"},
"tw-edu-parent-communication":{"recipients":"收件對象","purpose":"溝通目的","message":"訊息內容","requested_action":"期待配合事項","contact_channel":"聯絡管道"},
"tw-edu-pbl-designer":{"driving_question":"驅動問題","authentic_context":"真實情境","milestones":"里程碑","final_product":"最終成果","assessment_criteria":"評量準則"},
"tw-edu-school-document":{"document_type":"文件類型","basis":"依據","purpose":"目的","implementation":"實施方式","responsible_people":"權責人員","expected_results":"預期成果"},
"tw-edu-worksheet-creator":{"instructions":"作答說明","prompts":"學習任務","reflection":"反思問題"},
"tw-edu-anti-ai-assessment":{"items":"評量項目、向度分數與設計理由"},
}
FIELD_LABELS={"id":"編號","title":"名稱","instructions":"進行方式","materials":"材料","date":"日期","artifact":"作品","evidence":"證據","reflection":"反思","next_step":"下一步","situation":"情境","steps":"步驟","trigger":"觸發情況","teacher_action":"教師行動","follow_up":"後續追蹤","support":"支持方式","prompt":"題目","success_criteria":"成功準則","evidence_capture":"證據蒐集","action":"行動","discipline":"學科","knowledge":"知識內容","method":"方法","topic":"議題","minutes":"分鐘","owner":"負責人","due":"期限","deliverable":"交付成果","feedback":"回饋方式","item":"項目","details":"內容","response_space_lines":"作答行數","student_id":"學生代碼","observations":"具體觀察","strengths":"優勢","next_steps":"下一步","assessment_item":"評量題目","dimensions":"評分向度","total_score":"總分","reason":"理由","redesign":"調整方案","name":"名稱","score":"分數"}

def _strings(value: Any, prefix=""):
    if isinstance(value, dict):
        for k, v in value.items(): yield from _strings(v, f"{prefix}{k}｜")
    elif isinstance(value, list):
        for i, v in enumerate(value, 1): yield from _strings(v, f"{prefix}{i}｜")
    elif value is not None: yield prefix.rstrip("｜"), str(value)

def _docx(d: dict, output: Path, teacher=False) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    doc = Document(); sec = doc.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7); sec.top_margin=sec.bottom_margin=Cm(1.8); sec.left_margin=sec.right_margin=Cm(2)
    title = d["content"].get("title") or d["context"].get("topic") or d["skill"]
    doc.add_heading(title + ("｜教師答案" if teacher else ""), 0)
    doc.add_paragraph(f"{d['context'].get('subject','')}　{d['context'].get('grade','')}")
    doc.add_heading("來源與查核狀態",1)
    for source in d["sources"]: doc.add_paragraph(f"{source['title']}｜{source['status']}"+(f"｜{source['url']}" if source.get('url') else ""))
    if d["skill"] == "tw-edu-lesson-plan-108":
        c=d["content"]
        for heading,items,columns in [("學習目標",c["objectives"],["id","description"]),("教學活動",c["activities"],["id","title","minutes","instructions","objective_ids"]),("評量設計",c["assessments"],["id","method","criteria","objective_ids"])]:
            doc.add_heading(heading,1); table=doc.add_table(rows=1,cols=len(columns)); table.style="Table Grid"
            for j,key in enumerate(columns): table.rows[0].cells[j].text={"id":"編號","description":"目標敘述","title":"活動","minutes":"分鐘","instructions":"進行方式","objective_ids":"對應目標","method":"方法","criteria":"成功準則"}[key]
            for item in items:
                cells=table.add_row().cells
                for j,key in enumerate(columns): cells[j].text="、".join(item[key]) if isinstance(item[key],list) else str(item[key])
        if c.get("curriculum_codes"):
            doc.add_heading("課綱代碼與查核",1)
            for code in c["curriculum_codes"]: doc.add_paragraph(f"{code['code']}｜{'已查核' if code['verified'] else '待查核'}｜{code['description']}｜{code.get('verification_note','')}")
    elif d["skill"] == "tw-edu-rubric-designer":
        c=d["content"]; doc.add_heading("評量規準",1); levels=c["levels"]
        if c["type"]=="analytic":
            table=doc.add_table(rows=1,cols=2+len(levels)); table.style="Table Grid"
            heads=["向度","配分"]+[f"{x['label']}（{x['score']}）" for x in levels]
            for i,x in enumerate(heads): table.rows[0].cells[i].text=x
            for dim in c["dimensions"]:
                cells=table.add_row().cells; cells[0].text=dim["name"]; cells[1].text=str(dim["weight"])
                for i,lvl in enumerate(levels,2): cells[i].text=dim["descriptions"][lvl["id"]]
        else:
            table=doc.add_table(rows=1,cols=3); table.style="Table Grid"
            for i,x in enumerate(["等第","分數","整體描述"]): table.rows[0].cells[i].text=x
            for lvl in levels:
                cells=table.add_row().cells; cells[0].text=lvl["label"]; cells[1].text=str(lvl["score"]); cells[2].text=c["descriptions"][lvl["id"]]
        doc.add_paragraph(f"總分：{c['total_points']}")
    elif d["skill"] == "tw-edu-feedback-writer":
        for student in d["content"]["students"]:
            doc.add_heading(f"學生紀錄：{student['student_id']}",1)
            table=doc.add_table(rows=4,cols=2); table.style="Table Grid"
            for row,(label,key) in zip(table.rows,[("具體觀察","observations"),("優勢","strengths"),("下一步","next_steps"),("回饋文字","feedback")]): row.cells[0].text=label; row.cells[1].text="\n".join(student[key]) if isinstance(student[key],list) else student[key]
    elif d["skill"] == "tw-edu-learning-evidence-analyzer":
        c=d['content'];doc.add_paragraph(f"匿名學生數：{c['class_n']}")
        doc.add_heading('題目與分母',1);table=doc.add_table(rows=1,cols=7);table.style='Table Grid'
        for cell,label in zip(table.rows[0].cells,['題目／目標','已答','缺答','答對','全班正確率','已答正確率','答案狀態']): cell.text=label
        for item in c['items']:
            def percent(v): return '未計算' if v is None else f'{v:.1%}'
            values=[item['item_id']+'／'+item['target'],str(item['answered_n']),str(item['missing_n']),str(item['correct_n']) if item['correct_n'] is not None else '未計算',percent(item['accuracy_all_students']),percent(item['accuracy_answered']),{'verified':'已核實','pending':'待核實','ambiguous':'有疑義'}[item['key_status']]]
            for cell,text in zip(table.add_row().cells,values): cell.text=text
            doc.add_paragraph(item['item_id']+' 選項分布：'+ '、'.join(f'{key} {value}人' for key,value in item['option_counts'].items()))
        doc.add_heading('個別觀察與後續追問',1)
        for student,items in c['student_observations'].items():
            doc.add_paragraph(student+'：'+'；'.join(x['item_id']+' '+{'missing':'缺答','ungraded':'答案未核實','correct':'答對','needs_followup':'需追問原因'}[x['status']] for x in items))
        for note in c['limitations']: doc.add_paragraph(note)
        doc.add_paragraph('教學決策應另核學生解釋與作品；本表不自動判定迷思或介入效果。')
    elif d["skill"] == "tw-edu-material-reviewer":
        c=d['content'];doc.add_paragraph('檢查模式：'+{'material':'教材','assessment':'評量','slides':'簡報'}[c['mode']])
        for f in c['findings']:
            doc.add_heading(f['id']+'／'+f['material_id'],1)
            doc.add_paragraph('原文：'+f['quote']);doc.add_paragraph('狀態：'+{'confirmed':'已確認','unknown':'待查','resolved':'已修正'}[f['status']]+'｜影響：'+{'blocking':'影響使用','warning':'需調整','info':'建議','unknown':'待查'}[f['severity']]);doc.add_paragraph('理由：'+f['reason']);doc.add_paragraph('修正：'+f['revision'])
            if f['source_ids']: doc.add_paragraph('來源定位：'+ '、'.join(f['source_ids']))
        doc.add_heading('尚未檢查',1)
        for note in c['unchecked']: doc.add_paragraph(note)
        doc.add_paragraph('本報告編排使用者／分析者完成的檢查，不代表腳本自動核實所有知識或版面。')
    elif d["skill"] == "tw-edu-exam-generator":
        c=d["content"]
        doc.add_paragraph(f"題數：{len(c['questions'])} 題　總分：{c['total_points']} 分")
        for i,q in enumerate(c["questions"],1):
            doc.add_paragraph({'multiple_choice':'選擇題','short_answer':'簡答題','essay':'申論題','true_false':'是非題'}[q['type']])
            doc.add_heading(f"{i}. {q['prompt']}（{q['points']} 分）", 2)
            for o in q.get("options",[]): doc.add_paragraph(f"{o['id']}. {o['text']}")
            if teacher:
                doc.add_paragraph(f"答案：{q['answer']}")
                if q.get("explanation"): doc.add_paragraph(f"解析：{q['explanation']}")
            else: doc.add_paragraph("作答：________________________________")
    else:
        for key,val in d["content"].items():
            if key=="title": continue
            doc.add_heading(DOC_LABELS.get(d["skill"],{}).get(key,str(key)), 1)
            if isinstance(val,list):
                for item in val:
                    if isinstance(item,dict):
                        p=doc.add_paragraph(style="List Bullet")
                        p.add_run("；".join(f"{FIELD_LABELS.get(k,k)}：{v}" for k,v in item.items() if not isinstance(v,(list,dict))))
                        for k,v in item.items():
                            if isinstance(v,(list,dict)):
                                for label,text in _strings(v,f"{FIELD_LABELS.get(k,k)}｜"): doc.add_paragraph(f"{label}：{text}",style="List Bullet 2")
                        if d["skill"]=="tw-edu-worksheet-creator":
                            for _ in range(item.get("response_space_lines",0)): doc.add_paragraph("________________________________________________")
                    else: doc.add_paragraph(str(item),style="List Bullet")
            elif isinstance(val,dict):
                for label,text in _strings(val): doc.add_paragraph(f"{label}：{text}")
            else: doc.add_paragraph(str(val))
    styles=doc.styles
    for n in ["Normal","Title","Heading 1","Heading 2"]:
        styles[n].font.name="Noto Sans TC"; styles[n]._element.rPr.rFonts.set(qn("w:eastAsia"),"Noto Sans TC"); styles[n].font.size=Pt(11 if n=="Normal" else 16)
        styles[n].paragraph_format.line_spacing=1.35
        styles[n].font.color.rgb = RGBColor(0, 0, 0)
        for border in styles[n]._element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    for paragraph in doc.paragraphs:
        for border in paragraph._p.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    for table in doc.tables:
        header = OxmlElement('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(header)
        headings = [cell.text for cell in table.rows[0].cells]
        weights = [1 if h in {'編號','分鐘','配分'} else 2 if h == '對應目標' else 4 for h in headings]
        table.autofit = False
        for j, weight in enumerate(weights):
            table.columns[j].width = Cm(17 * weight / sum(weights))
        for row in table.rows:
            for j, cell in enumerate(row.cells):
                cell.width = Cm(17 * weights[j] / sum(weights))
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.line_spacing = 1.15
                    for run in p.runs:
                        run.font.name = 'Noto Sans TC'
                        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans TC')
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name="Noto Sans TC"; run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Noto Sans TC")
    output.parent.mkdir(parents=True,exist_ok=True); doc.save(output)

def _xlsx(d: dict, output: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    c=d["content"]; wb=Workbook(); ws=wb.active; ws.title="課程地圖"
    ws.append([c["title"]]); ws.merge_cells("A1:F1"); ws["A1"].font=Font(bold=True,size=16)
    ws.append(["科目",d["context"]["subject"],"年級",d["context"]["grade"],"主題",d["context"]["topic"]])
    ws.append(["來源", "；".join(f"{x['title']}({x['status']})" for x in d["sources"])])
    headers=["單元","週次","節數","目標","課綱代碼","評量"]
    ws.append(headers)
    for cell in ws[4]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="1A5276")
    for u in c["units"]: ws.append([u["name"],u["weeks"],u["periods"],"\n".join(u["goals"]),"\n".join(f"{x['code']}｜{'已查核' if x['verified'] else '待查核'}｜{x.get('verification_note','')}" for x in u["codes"]),"\n".join(u["assessments"])])
    for col,w in zip("ABCDEF",[22,12,8,40,22,35]): ws.column_dimensions[col].width=w
    for row in ws.iter_rows():
        for cell in row: cell.alignment=Alignment(vertical="top",wrap_text=True)
    ws.freeze_panes = 'A5'
    ws.print_title_rows = '1:4'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    output.parent.mkdir(parents=True,exist_ok=True); wb.save(output)

def _safe_json(v): return json.dumps(v,ensure_ascii=False).replace("<","\\u003c")

def _html(d: dict, output: Path) -> None:
    c=d["content"]; mode=c["mode"]; payload=_safe_json(c)
    script = Path(__file__).with_name('miniapp.js').read_text(encoding='utf-8')
    title=html.escape(c["title"])
    page=f'''<!doctype html><html lang="zh-Hant"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>{title}</title><style>body{{font-family:system-ui,sans-serif;max-width:56rem;margin:auto;padding:2rem;background:#f7fafc;color:#17202a}}button{{display:block;margin:.6rem 0;padding:.8rem 1rem}}button:focus{{outline:3px solid #2471a3}}</style><h1>{title}</h1><p>{html.escape(d['context']['subject'])}｜{html.escape(d['context']['grade'])}</p><main id="app"></main><script id="data" type="application/json">{payload}</script><script>{script}</script></html>'''
    output.parent.mkdir(parents=True,exist_ok=True); output.write_text(page,encoding="utf-8")

def _font() -> str:
    roots=[Path("/System/Library/Fonts"),Path("/Library/Fonts"),Path("/usr/share/fonts"),Path.home()/"Library/Fonts"]
    preferred=("NotoSansCJK","NotoSansTC","PingFang","JhengHei","Songti")
    for root in roots:
        if root.exists():
            files=[p for p in root.rglob("*") if p.suffix.lower() in {".ttf",".otf",".ttc"}]
            for token in preferred:
                for path in files:
                    if token.lower() in path.name.lower(): return str(path)
    raise InputError("CJK font unavailable; install Noto Sans CJK or another Traditional Chinese font")

def _png(d: dict, output: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont
    c=d["content"]; font_path=_font(); title_font=ImageFont.truetype(font_path,34); body_font=ImageFont.truetype(font_path,24)
    if c["type"]=="prisma":
        p=c["prisma"]; labels=[("辨識（去除重複 "+str(p["duplicates_removed"])+"）",p["identified"]),("篩選（排除 "+str(p["screening_excluded"])+"）",p["screened"]),("全文審查（排除 "+str(p["full_text_excluded"])+"）",p["full_text_assessed"]),("納入",p["included"])]
    else: labels=[(x["label"],x.get("value","")) for x in c["nodes"]]
    image=Image.new("RGB",(1600,900),"white"); draw=ImageDraw.Draw(image)
    draw.text((800,45),c["title"]+"（簡易流程圖）",font=title_font,fill="#1A5276",anchor="ma")
    gap=700/max(1,len(labels)); box_h=min(105,gap*.7)
    for i,(label,value) in enumerate(labels):
        y=125+i*gap; draw.rounded_rectangle((480,y,1120,y+box_h),radius=16,fill="#EBF5FB",outline="#2471A3",width=3)
        draw.multiline_text((800,y+box_h/2),f"{label}\n{value}",font=body_font,fill="#17202A",anchor="mm",align="center")
        if i<len(labels)-1: draw.line((800,y+box_h,800,y+gap),fill="#2471A3",width=4); draw.polygon([(790,y+gap-12),(810,y+gap-12),(800,y+gap)],fill="#2471A3")
    output.parent.mkdir(parents=True,exist_ok=True); image.save(output)

def _report(input_path: Path, outputs: list[Path], d: dict, sample: bool) -> Path:
    primary=outputs[0]; report=primary.with_name(primary.name+".validation.json")
    checks=[]
    if any(x["status"]!="verified" for x in d["sources"]): checks.append("source_verification")
    if d["skill"]=="tw-edu-lesson-plan-108" and any(not x.get("verified") for x in d["content"].get("curriculum_codes",[])): checks.append("curriculum_code_verification")
    if d["skill"]=="tw-edu-curriculum-mapper" and any(not x.get("verified") for u in d["content"]["units"] for x in u["codes"]): checks.append("curriculum_code_verification")
    checks.append("human_visual_review")
    body={"schema_version":"1.0","skill":d["skill"],"sample":sample,"input":{"path":str(input_path),"sha256":_hash(input_path)},"outputs":[{"path":str(x),"sha256":_hash(x)} for x in outputs],"pending_checks":checks}
    report.write_text(json.dumps(body,ensure_ascii=False,indent=2)+"\n",encoding="utf-8"); return report

def run(skill: str, input_path: Path, output: Path, validate_only=False, sample=False) -> list[Path]:
    root=_skill_dir(); schema=root/"schemas"/"input.schema.json"; d=_json(input_path)
    _validate(d,schema)
    if d.get("skill") != skill: raise InputError(f"input skill must be {skill}")
    _semantic(skill,d)
    if skill == 'tw-edu-slides-creator':
        from .slides import prepare
        try: prepare(d['content'], input_path.resolve().parent)
        except (ValueError, OSError) as exc: raise InputError(str(exc)) from exc
    if validate_only: return []
    if sample and '範例' not in d['content']['title']:
        d['content']['title'] += '（範例）'
    if output.exists(): raise InputError(f'output already exists: {output}')
    kind=SKILLS[skill]; outputs=[]
    if skill == 'tw-edu-learning-evidence-analyzer':
        from .learning_evidence import analyze
        summary=analyze(d['content'])
        output.parent.mkdir(parents=True,exist_ok=True)
        evidence_file=output.with_suffix('.analysis.json')
        if evidence_file.exists(): raise InputError('analysis output already exists')
        evidence_file.write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n')
        import copy
        d=copy.deepcopy(d);d['content']={'title':d['content']['title'],**summary}

    if kind=="exam":
        stem=output.with_suffix(""); student=stem.with_name(stem.name+"-student").with_suffix(".docx"); teacher=stem.with_name(stem.name+"-teacher").with_suffix(".docx")
        if student.exists() or teacher.exists(): raise InputError('exam output already exists')
        _docx(d,student,False); _docx(d,teacher,True); outputs=[student,teacher]
    else:
        expected={"docx":".docx","xlsx":".xlsx","html":".html","png":".png","pptx":".pptx"}[kind]
        if output.suffix.lower()!=expected: raise InputError(f"output must use {expected}")
        from .slides import render as render_slides
        {"docx":_docx,"xlsx":_xlsx,"html":_html,"png":_png,"pptx":render_slides}[kind](d,output); outputs=[output]
    if skill == "tw-edu-learning-evidence-analyzer": outputs.append(evidence_file)
    _report(input_path,outputs,d,sample); return outputs

def main(skill_name: str) -> int:
    if skill_name not in SKILLS: print(f"unsupported skill: {skill_name}",file=sys.stderr); return 2
    p=argparse.ArgumentParser(description=f"Input-driven generator for {skill_name}")
    p.add_argument("--input",type=Path); p.add_argument("--output",type=Path); p.add_argument("--validate-only",action="store_true"); p.add_argument("--example",action="store_true")
    args,unknown=p.parse_known_args()
    if unknown: p.error("legacy/unknown flags are unsupported; migrate to --input, --output, --validate-only, or --example")
    if args.example and args.input: p.error("--example and --input are mutually exclusive")
    if not args.example and not args.input: p.error("--input is required unless --example is used")
    if not args.validate_only and not args.output:
        extension = 'docx' if SKILLS[skill_name] == 'exam' else SKILLS[skill_name]
        args.output = Path.cwd() / 'artifacts' / skill_name / uuid.uuid4().hex[:12] / ('output.' + extension)
    source=_skill_dir()/"examples"/"example.json" if args.example else args.input
    try:
        outputs=run(skill_name,source,args.output,args.validate_only,args.example)
        print("valid" if args.validate_only else "generated: "+", ".join(str(x) for x in outputs)); return 0
    except (InputError, OSError, ValueError) as e: print(f"error: {e}",file=sys.stderr); return 2
