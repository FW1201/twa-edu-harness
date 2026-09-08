#!/usr/bin/env python3
"""擷取 .docx 的版面指紋，供重構前後比對。

只比表格數與段落數是不夠的——把行內實作換成共用版時，真正會壞掉的是
儲存格底色、框線色、字型、對齊這些細節，而它們在檔案大小上看不出差別。
教師手上已經有用這些技能產出的檔案，版面不能悄悄改變。

用法：
    python scripts/docx_fingerprint.py <file.docx>            # 輸出 JSON 指紋
    python scripts/docx_fingerprint.py --diff <a.docx> <b.docx>
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def _run_style(run) -> dict:
    rPr = run._r.find(qn("w:rPr"))
    east_asia = None
    if rPr is not None:
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is not None:
            east_asia = rFonts.get(qn("w:eastAsia"))
    return {
        "text": run.text,
        "bold": bool(run.bold),
        "size": run.font.size.pt if run.font.size else None,
        "name": run.font.name,
        "eastAsia": east_asia,
        "color": str(run.font.color.rgb) if run.font.color and run.font.color.rgb else None,
    }


def _para(p) -> dict:
    return {
        "text": p.text,
        "alignment": str(p.alignment) if p.alignment is not None else None,
        "runs": [_run_style(r) for r in p.runs],
        "hasBottomBorder": p._p.find(qn("w:pPr")) is not None
        and p._p.find(qn("w:pPr")).find(qn("w:pBdr")) is not None,
    }


def _cell(c) -> dict:
    """儲存格指紋。

    ⚠️ 必須讀**所有** w:shd 與 w:tcBorders，不能只讀第一個。
    `set_cell_bg()` 是 append 而非取代，重複套用會留下多個 w:shd；
    只讀第一個的話，底色改變會偵測不到——這個工具就失去意義了。
    """
    tcPr = c._tc.find(qn("w:tcPr"))
    shadings: list[str | None] = []
    border_sets: list[dict] = []
    if tcPr is not None:
        for shd in tcPr.findall(qn("w:shd")):
            shadings.append(shd.get(qn("w:fill")))
        for bd in tcPr.findall(qn("w:tcBorders")):
            border_sets.append({
                side: (e.get(qn("w:color")), e.get(qn("w:sz")))
                for side in ("top", "bottom", "left", "right")
                if (e := bd.find(qn(f"w:{side}"))) is not None
            })
    return {"shadings": shadings, "borderSets": border_sets,
            "paragraphs": [_para(p) for p in c.paragraphs]}


def fingerprint(path: Path) -> dict:
    doc = Document(str(path))
    sec = doc.sections[0]
    return {
        "section": {
            "pageWidth": sec.page_width, "pageHeight": sec.page_height,
            "margins": [sec.top_margin, sec.bottom_margin,
                        sec.left_margin, sec.right_margin],
        },
        "normalStyle": {
            "font": doc.styles["Normal"].font.name,
            "size": doc.styles["Normal"].font.size.pt
            if doc.styles["Normal"].font.size else None,
        },
        "paragraphs": [_para(p) for p in doc.paragraphs],
        "tables": [
            {"rows": len(t.rows), "cols": len(t.columns), "style": t.style.name,
             "cells": [[_cell(c) for c in row.cells] for row in t.rows]}
            for t in doc.tables
        ],
        "headerText": doc.sections[0].header.paragraphs[0].text,
    }


def _walk(a, b, path=""):
    """回傳所有差異點，路徑化以便定位。"""
    if type(a) is not type(b):
        return [f"{path}: 型別 {type(a).__name__} → {type(b).__name__}"]
    if isinstance(a, dict):
        out = []
        for k in sorted(set(a) | set(b)):
            if k not in a:
                out.append(f"{path}.{k}: 新增 {b[k]!r}")
            elif k not in b:
                out.append(f"{path}.{k}: 消失 {a[k]!r}")
            else:
                out.extend(_walk(a[k], b[k], f"{path}.{k}"))
        return out
    if isinstance(a, list):
        out = []
        if len(a) != len(b):
            out.append(f"{path}: 長度 {len(a)} → {len(b)}")
        for i, (x, y) in enumerate(zip(a, b)):
            out.extend(_walk(x, y, f"{path}[{i}]"))
        return out
    return [] if a == b else [f"{path}: {a!r} → {b!r}"]


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+")
    ap.add_argument("--diff", action="store_true")
    ap.add_argument("--max", type=int, default=30, help="最多列出幾項差異")
    args = ap.parse_args()

    if args.diff:
        if len(args.files) != 2:
            print("❌ --diff 需要兩個檔案")
            return 1
        a, b = (fingerprint(Path(f)) for f in args.files)
        diffs = _walk(a, b, "doc")
        if not diffs:
            print("✅ 版面指紋完全相同")
            return 0
        print(f"❌ {len(diffs)} 項版面差異：")
        for d in diffs[:args.max]:
            print(f"  - {d}")
        if len(diffs) > args.max:
            print(f"  …另有 {len(diffs) - args.max} 項")
        return 1

    print(json.dumps(fingerprint(Path(args.files[0])),
                     ensure_ascii=False, indent=1, default=str))
    return 0


if __name__ == "__main__":
    sys.exit(main())
